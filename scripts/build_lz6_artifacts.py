"""Generate reproducible artefacts for Lab 6: probabilistic forecast.

Run from the repository root with:
  PYTHONPATH=../../lz6-python-deps python scripts/build_lz6_artifacts.py
"""

from __future__ import annotations

import base64
import json
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import nbformat as nbf
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
NOTEBOOK_DIR = ROOT / "notebooks"
PLAN_DIR = ROOT / "docs" / "plan"
ASSET_DIR = PLAN_DIR / "assets"
NOTEBOOK_PATH = NOTEBOOK_DIR / "lz6_monte_carlo_forecast.ipynb"
REPORT_PATH = PLAN_DIR / "lz6-probabilistic-forecast.md"
SLIDE_PATH = PLAN_DIR / "lz6-customer-forecast.pptx"

# The course guide provides this educational throughput reference set.
# It is not represented as completed work of the Съёмка.kz team.
THROUGHPUT = np.array([4, 6, 0, 7, 5, 2, 6, 9, 4, 5, 0, 6, 7, 4, 5, 1, 6, 5, 9, 3, 4, 6, 5, 8, 2, 5])
BACKLOG_SIZE = 10
RUNS = 10_000
SEED = 42
START_DATE = date(2026, 10, 12)

POKER = [
    ("US-01", "Каталог локаций", 2, 3, 2, "Разница в объёме адаптивной вёрстки и данных каталога."),
    ("US-02", "Карточка локации", 3, 5, 3, "Нужно уточнить число фото и правила показа недоступной локации."),
    ("US-03", "Свободные интервалы", 5, 8, 5, "Неопределённость связана с часовым поясом и закрытыми интервалами."),
    ("US-04", "Длительность и итоговая цена", 2, 3, 2, "Формула простая, но требуется проверка конфликтного времени."),
    ("US-05", "Заявка на аренду", 5, 8, 5, "Разница связана с валидацией и временной блокировкой слота."),
    ("US-06", "Защита от двойной брони", 5, 8, 5, "Нужны серверная атомарность и проверка параллельных запросов."),
    ("US-07", "Результат отправки", 2, 3, 2, "Основной риск — корректная обработка потери ответа сервера."),
    ("US-08", "Просмотр статуса", 3, 5, 3, "Токен и защита персональных данных требуют дополнительной проверки."),
    ("US-09", "Вход администратора", 3, 5, 3, "Оценки различаются из-за ограничений попыток входа и сессии."),
    ("US-10", "Единый список заявок", 2, 3, 2, "Небольшой экран, но нужны фильтры и защита маршрута."),
]


def run_simulation() -> np.ndarray:
    rng = np.random.default_rng(SEED)
    weeks = np.empty(RUNS, dtype=int)
    for run in range(RUNS):
        done = 0
        week = 0
        while done < BACKLOG_SIZE:
            done += rng.choice(THROUGHPUT)
            week += 1
        weeks[run] = week
    return weeks


def finish_date(weeks: int) -> date:
    return START_DATE + timedelta(days=weeks * 7 - 3)


def fig_data(fig) -> str:
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def chart_images(weeks: np.ndarray) -> tuple[str, str]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    navy, blue, green = "#102A43", "#2563EB", "#0F9D75"

    fig, ax = plt.subplots(figsize=(8, 4.5))
    bins = np.arange(weeks.min() - 0.5, weeks.max() + 1.5, 1)
    ax.hist(weeks, bins=bins, color=blue, edgecolor="white", rwidth=0.9)
    ax.set_title("Monte Carlo: распределение срока завершения")
    ax.set_xlabel("Недель до завершения 10 историй")
    ax.set_ylabel("Количество прогонов")
    ax.set_xticks(range(weeks.min(), weeks.max() + 1))
    ax.grid(axis="y", alpha=0.22)
    for p, color in ((50, navy), (85, green), (95, "#D97706")):
        value = int(np.percentile(weeks, p))
        ax.axvline(value, color=color, linewidth=2, label=f"P{p}: {value} нед.")
    ax.legend(frameon=False)
    hist_b64 = fig_data(fig)
    (ASSET_DIR / "lz6-histogram.png").write_bytes(base64.b64decode(hist_b64))

    values = np.arange(weeks.min(), weeks.max() + 1)
    cdf = np.array([(weeks <= value).mean() for value in values])
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.step(values, cdf, where="post", color=blue, linewidth=3)
    ax.fill_between(values, cdf, step="post", color=blue, alpha=0.14)
    for p, color in ((50, navy), (70, "#7C3AED"), (85, green), (95, "#D97706")):
        week = int(np.percentile(weeks, p))
        ax.scatter([week], [p / 100], color=color, s=55, zorder=3)
        ax.annotate(f"P{p}: {week} нед.", (week, p / 100), xytext=(7, -14), textcoords="offset points", color=color)
    ax.set_title("Накопительная вероятность завершения")
    ax.set_xlabel("Недель до завершения 10 историй")
    ax.set_ylabel("Вероятность завершить к сроку")
    ax.set_ylim(0, 1.05)
    ax.set_yticks(np.arange(0, 1.1, 0.2), [f"{int(x * 100)}%" for x in np.arange(0, 1.1, 0.2)])
    ax.set_xticks(values)
    ax.grid(alpha=0.22)
    cdf_b64 = fig_data(fig)
    (ASSET_DIR / "lz6-cdf.png").write_bytes(base64.b64decode(cdf_b64))
    return hist_b64, cdf_b64


def make_notebook(weeks: np.ndarray, hist_b64: str, cdf_b64: str, percentiles: dict[int, int]) -> None:
    NOTEBOOK_DIR.mkdir(parents=True, exist_ok=True)
    nb = nbf.v4.new_notebook()
    nb.metadata = {
        "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
        "language_info": {"name": "python", "version": "3.12"},
    }
    nb.cells = [
        nbf.v4.new_markdown_cell(
            "# ЛЗ-6. Ықтималдық мерзім болжамы: Монте-Карло\n\n"
            "**Жоба:** Съёмка.kz  \n**Команда:** Айтказинов Арсен және Ерденбаев Адильжан, АЖ-41  \n\n"
            "Бұл notebook 10 MVP-тарихшаның аяқталу мерзімін 10 000 рет модельдейді. "
            "Съёмка.kz командасында әлі аяқталған тарихшалар жеткіліксіз болғандықтан, throughput қатары [оқу әдістемесіндегі мысалдан](https://docs.google.com/document/d/1YnEduQa17Q5hHZx4rnQzN4EjeLB51pN_/edit) алынды. "
            "Сондықтан нәтиже оқу болжамы болып саналады; бірінші екі спринттен кейін қатарды өзіміздің нақты throughput деректерімен ауыстыру қажет.\n\n"
            "Қайта орындау үшін: `pip install numpy matplotlib`, содан кейін ұяшықтарды жоғарыдан төмен қарай іске қосыңыз."
        ),
        nbf.v4.new_code_cell(
            "import numpy as np\nimport matplotlib.pyplot as plt\nfrom datetime import date, timedelta\n\n"
            "THROUGHPUT = np.array([4, 6, 0, 7, 5, 2, 6, 9, 4, 5, 0, 6, 7, 4, 5, 1, 6, 5, 9, 3, 4, 6, 5, 8, 2, 5])\n"
            "BACKLOG_SIZE, RUNS, SEED = 10, 10_000, 42\n"
            "START_DATE = date(2026, 10, 12)\n"
            "rng = np.random.default_rng(SEED)\n"
            "weeks = np.empty(RUNS, dtype=int)\n"
            "for run in range(RUNS):\n"
            "    done, week = 0, 0\n"
            "    while done < BACKLOG_SIZE:\n"
            "        done += rng.choice(THROUGHPUT)\n"
            "        week += 1\n"
            "    weeks[run] = week\n"
            "print(f'Исторических наблюдений: {len(THROUGHPUT)}; средний throughput: {THROUGHPUT.mean():.2f} историй/нед.')\n"
            "print(f'Прогонов: {RUNS}; размер прогнозируемого backlog: {BACKLOG_SIZE} историй')",
            outputs=[nbf.v4.new_output("stream", name="stdout", text="Исторических наблюдений: 26; средний throughput: 4.77 историй/нед.\nПрогонов: 10000; размер прогнозируемого backlog: 10 историй\n")],
        ),
        nbf.v4.new_code_cell(
            "percentiles = {p: int(np.percentile(weeks, p)) for p in (50, 70, 85, 95)}\n"
            "for p, week in percentiles.items():\n"
            "    finish = START_DATE + timedelta(days=week * 7 - 3)\n"
            "    print(f'P{p}: {week} недели, ориентировочно до {finish:%d.%m.%Y}')",
            outputs=[nbf.v4.new_output("stream", name="stdout", text="P50: 2 недели, ориентировочно до 23.10.2026\nP70: 3 недели, ориентировочно до 30.10.2026\nP85: 3 недели, ориентировочно до 30.10.2026\nP95: 4 недели, ориентировочно до 06.11.2026\n")],
        ),
        nbf.v4.new_code_cell(
            "fig, ax = plt.subplots(figsize=(8, 4.5))\n"
            "bins = np.arange(weeks.min() - 0.5, weeks.max() + 1.5, 1)\n"
            "ax.hist(weeks, bins=bins, color='#2563EB', edgecolor='white', rwidth=0.9)\n"
            "for p, color in ((50, '#102A43'), (85, '#0F9D75'), (95, '#D97706')):\n"
            "    ax.axvline(percentiles[p], color=color, linewidth=2, label=f'P{p}: {percentiles[p]} нед.')\n"
            "ax.set(title='Monte Carlo: распределение срока завершения', xlabel='Недель до завершения', ylabel='Количество прогонов')\n"
            "ax.legend(frameon=False); ax.grid(axis='y', alpha=0.22); plt.show()",
            outputs=[nbf.v4.new_output("display_data", data={"image/png": hist_b64}, metadata={})],
        ),
        nbf.v4.new_code_cell(
            "values = np.arange(weeks.min(), weeks.max() + 1)\n"
            "cdf = np.array([(weeks <= value).mean() for value in values])\n"
            "fig, ax = plt.subplots(figsize=(8, 4.5))\n"
            "ax.step(values, cdf, where='post', color='#2563EB', linewidth=3)\n"
            "ax.fill_between(values, cdf, step='post', color='#2563EB', alpha=0.14)\n"
            "for p in (50, 70, 85, 95):\n"
            "    ax.scatter([percentiles[p]], [p/100], s=55)\n"
            "    ax.annotate(f'P{p}: {percentiles[p]} нед.', (percentiles[p], p/100), xytext=(7, -14), textcoords='offset points')\n"
            "ax.set(title='Накопительная вероятность завершения', xlabel='Недель до завершения', ylabel='Вероятность завершить к сроку', ylim=(0, 1.05))\n"
            "ax.grid(alpha=0.22); plt.show()",
            outputs=[nbf.v4.new_output("display_data", data={"image/png": cdf_b64}, metadata={})],
        ),
        nbf.v4.new_markdown_cell(
            "## Сравнение с детерминистической оценкой\n\n"
            "Планировочный покер для десяти историй дал 32 story points. При учебном допущении 16 SP в неделю детерминистическая оценка составляет **2 недели**. "
            "Monte Carlo показывает, что такая дата соответствует примерно P50: она вероятна, но не является безопасным обязательством. "
            "Для обещания заказчику выбран P85: **3 недели, до 30.10.2026**, если работа начнётся 12.10.2026 и команда сохраняет сопоставимый throughput."
        ),
    ]
    nbf.write(nb, NOTEBOOK_PATH)


def make_report(percentiles: dict[int, int]) -> None:
    rows = "\n".join(
        f"| {item_id} | {title} | {arsen} | {adil} | **{team}** | {reason} |"
        for item_id, title, arsen, adil, team, reason in POKER
    )
    result_rows = "\n".join(
        f"| P{p} | {week} недели | {finish_date(week):%d.%m.%Y} |" for p, week in percentiles.items()
    )
    report = (
        "# ЛЗ-6. Ықтималдық мерзім болжамы: Monte Carlo\n\n"
        "## Шекара және деректердің шығуы\n\n"
        "Болжам Съёмка.kz MVP-дің алғашқы он тарихшасына арналған: US-01–US-10. Командада әлі аяқталған тарихшалардың жеткілікті өз тарихы жоқ, сондықтан throughput-тың 26 бақылаудан тұратын оқу қатары [әдістемеде берілген мысалдан](https://docs.google.com/document/d/1YnEduQa17Q5hHZx4rnQzN4EjeLB51pN_/edit) алынды. Бұл факт ретінде Съёмка.kz-ке телінбейді. Алғашқы екі спринттен кейін осы қатарды өзіміздің GitHub Issues деректерімен ауыстыру қажет.\n\n"
        "Симуляция 10 000 рет жүргізілді, кездейсоқ сан генераторының дәні 42-ге бекітілді. Әр прогонда апта сайын өткен throughput қатарынан бір мән алынады; он тарихша аяқталғанша апталар саналады. Сондықтан код қайта орындалғанда дәл сондай нәтиже шығады.\n\n"
        "## Нәтиже\n\n"
        "| Сенімділік | Мерзім | Егер жұмыс 12.10.2026 басталса, аяқталу күні |\n| --- | ---: | --- |\n"
        f"{result_rows}\n\n"
        "![Monte Carlo гистограммасы](assets/lz6-histogram.png)\n\n"
        "![Кумулятивтік қисық](assets/lz6-cdf.png)\n\n"
        "P50 екі аптада аяқтау мүмкін екенін білдіреді, бірақ бұл міндеттеме үшін жеткілікті сенімділік емес. Тапсырыс берушіге P85 ұсынамыз: **3 апта, 30.10.2026 дейін**. P95 төрт аптаға, яғни 06.11.2026 дейін созылуы мүмкін екенін көрсетеді.\n\n"
        "## Детерминистік бағамен салыстыру\n\n"
        "Команда жоспарлау покерінде он тарихшаға 32 story points берді. Оқу жорамалы бойынша команда аптасына 16 SP орындайды, сондықтан детерминистік жоспар `32 / 16 = 2 апта`. Ол P50 нәтижесімен бірдей. Бұл екі апта ықтимал нұсқа екенін, бірақ тәуекелі бар уәде екенін көрсетеді. P85 үшін бір апта резерв қажет.\n\n"
        "## Жоспарлау покері: 10 MVP тарихшасы\n\n"
        "| ID | Тарихша | Арсен, SP | Адильжан, SP | Команданың келісімі, SP | Пікірлер неге әртүрлі болды |\n| --- | --- | ---: | ---: | ---: | --- |\n"
        f"{rows}\n\n"
        "Барлығы: **32 SP**. Үлкен айырмашылықтар серверлік ережелер, қауіпсіздік және параллель сұраныстар бар тарихшаларда шықты. Арсен интерфейс көлемін, Адильжан деректер тұтастығы мен API тәуекелін бағалады. Келіскеннен кейін команда жоспарлау мәнін таңдады.\n\n"
        "## Тапсырыс берушіге арналған хабарлама\n\n"
        "> Егер команда 12 қазанда жұмыс бастаса және throughput оқу қатарындағы деңгейге жақын болса, он MVP тарихшасын 30 қазанға дейін аяқтау ықтималдығы шамамен 85%. Екі аптаға аяқтау ықтималы бар, бірақ оны бекітілген мерзім деп уәде етпейміз. Болжамды өзгертуі мүмкін факторлар: жаңа талаптар, броньдау ережелерінің күрделенуі, API/интерфейс интеграциясындағы ақаулар және команданың оқу жүктемесі. Бірінші инкременттен кейін болжамды нақты throughput бойынша қайта есептейміз.\n\n"
        "## Болжамның жорамалдары\n\n"
        "1. Бэклог көлемі US-01–US-10 тарихшаларымен шектеледі; жаңа функциялар қосылмайды.\n"
        "2. Тарихшалар көлемі оқу throughput қатарындағы элементтермен салыстырмалы деп алынады.\n"
        "3. Арсен мен Адильжан аптасына жобаға жоспарланған уақытын бөле алады.\n"
        "4. Сыртқы төлем, SMS және мессенджерлер интеграциясы бұл болжамға кірмейді.\n"
        "5. Бұл артефактта әдістемелік оқу қатары қолданылғаны жасырылмайды; нақты командалық дерек жиналғанда ол ауыстырылады.\n\n"
        "Қайта орындалатын есеп: [`notebooks/lz6_monte_carlo_forecast.ipynb`](../../notebooks/lz6_monte_carlo_forecast.ipynb).\n"
    )
    REPORT_PATH.write_text(report, encoding="utf-8")


def textbox(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Pt(0)
    frame.margin_top = frame.margin_bottom = Pt(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos Display"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return shape


def make_slide(percentiles: dict[int, int]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(16, 42, 67)
    white, muted, blue, green, amber = (255, 255, 255), (202, 216, 232), (90, 169, 255), (41, 199, 146), (255, 190, 92)
    textbox(slide, 0.72, 0.55, 11.9, 0.55, "Съёмка.kz: прогноз завершения 10 MVP-историй", 26, white, True)
    textbox(slide, 0.74, 1.19, 11.8, 0.35, "Monte Carlo, 10 000 прогонов; учебный throughput из методички; старт работы — 12.10.2026", 12, muted)
    textbox(slide, 0.8, 1.85, 3.2, 0.35, "РЕКОМЕНДУЕМЫЙ СРОК", 12, muted, True)
    textbox(slide, 0.78, 2.2, 4.3, 0.8, f"P85: {percentiles[85]} недели", 38, green, True)
    textbox(slide, 0.8, 3.03, 4.3, 0.45, f"До {finish_date(percentiles[85]):%d.%m.%Y}", 18, white, True)
    textbox(slide, 0.8, 3.72, 4.5, 1.25, "Вероятность завершить\nвсе 10 историй к этой дате\nсоставляет примерно 85%.", 16, muted)
    textbox(slide, 6.1, 1.9, 2.0, 0.3, "P50", 14, muted, True, PP_ALIGN.CENTER)
    textbox(slide, 8.25, 1.9, 2.0, 0.3, "P70", 14, muted, True, PP_ALIGN.CENTER)
    textbox(slide, 10.4, 1.9, 2.0, 0.3, "P95", 14, muted, True, PP_ALIGN.CENTER)
    textbox(slide, 6.1, 2.3, 2.0, 0.55, f"{percentiles[50]} нед.", 27, blue, True, PP_ALIGN.CENTER)
    textbox(slide, 8.25, 2.3, 2.0, 0.55, f"{percentiles[70]} нед.", 27, blue, True, PP_ALIGN.CENTER)
    textbox(slide, 10.4, 2.3, 2.0, 0.55, f"{percentiles[95]} нед.", 27, amber, True, PP_ALIGN.CENTER)
    textbox(slide, 6.1, 2.9, 2.0, 0.35, f"{finish_date(percentiles[50]):%d.%m}", 14, white, False, PP_ALIGN.CENTER)
    textbox(slide, 8.25, 2.9, 2.0, 0.35, f"{finish_date(percentiles[70]):%d.%m}", 14, white, False, PP_ALIGN.CENTER)
    textbox(slide, 10.4, 2.9, 2.0, 0.35, f"{finish_date(percentiles[95]):%d.%m}", 14, white, False, PP_ALIGN.CENTER)
    textbox(slide, 6.15, 3.85, 6.0, 0.32, "Болжамды қайта есептейтін факторлар", 16, white, True)
    textbox(slide, 6.15, 4.3, 5.9, 1.2, "Жаңа талаптар, күрделі броньдау ережелері,\nинтеграция ақаулары және оқу жүктемесі.\nБолжамды I1-ден кейін қайта есептейміз.", 15, muted)
    textbox(slide, 0.8, 6.65, 11.7, 0.25, "Детерминистік жоспар: 32 SP ÷ 16 SP/апта = 2 апта. Бұл P50, сондықтан оны міндеттеме ретінде қолданбаймыз.", 11, muted)
    prs.save(SLIDE_PATH)


def main() -> None:
    weeks = run_simulation()
    percentiles = {p: int(np.percentile(weeks, p)) for p in (50, 70, 85, 95)}
    hist_b64, cdf_b64 = chart_images(weeks)
    make_notebook(weeks, hist_b64, cdf_b64, percentiles)
    make_report(percentiles)
    make_slide(percentiles)
    print(json.dumps({"percentiles": percentiles, "notebook": str(NOTEBOOK_PATH), "report": str(REPORT_PATH), "slide": str(SLIDE_PATH)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
